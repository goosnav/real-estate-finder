"""Bounded document imports and spreadsheet-safe exports."""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pypdf import PdfReader

from .config import AppError
from .quality import client_signature

MAX_FILE = 8 * 1024 * 1024
MAX_ROWS = 10000


def check_file(raw: bytes):
    if not raw:
        raise AppError("FILE-EMPTY", "Choose a file with content.")
    if len(raw) > MAX_FILE:
        raise AppError("FILE-SIZE", "Files must be 8 MB or smaller.", 413)


def check_zip(raw: bytes):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            items = archive.infolist()
            if len(items) > 2000 or sum(i.file_size for i in items) > 40 * 1024 * 1024:
                raise AppError("FILE-SIZE", "The expanded document is too large.")
            if any(i.flag_bits & 1 for i in items):
                raise AppError("FILE-ENCRYPTED", "Use an unencrypted document.")
    except zipfile.BadZipFile as exc:
        raise AppError("FILE-FORMAT", "This document is not a valid Office file.") from exc


def document_text(name: str, raw: bytes) -> str:
    check_file(raw)
    suffix = Path(name).suffix.lower()
    try:
        if suffix in {".txt", ".md", ".csv"}:
            text = raw.decode("utf-8-sig")
        elif suffix == ".pdf":
            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted:
                raise AppError("FILE-ENCRYPTED", "Use an unencrypted PDF.")
            if len(reader.pages) > 80:
                raise AppError("FILE-SIZE", "Use a PDF with 80 pages or fewer.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx":
            check_zip(raw)
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                xml = archive.read("word/document.xml")
            root = ElementTree.fromstring(xml)
            text = "\n".join(
                "".join(p.itertext())
                for p in root.iter(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"
                )
            )
        else:
            raise AppError("FILE-TYPE", "Use TXT, Markdown, PDF, DOCX, or CSV for your brief or supporting material.")
    except AppError:
        raise
    except Exception as exc:
        raise AppError(
            "FILE-READ", "The document could not be read. Try a plain-text export."
        ) from exc
    text = text.replace("\x00", "").strip()
    if not text:
        raise AppError("FILE-NO-TEXT", "No text found. Scanned PDFs need OCR before upload.")
    if len(text) > 60000:
        raise AppError("FILE-SIZE", "Use a shorter document (up to 60,000 text characters).")
    return text


def header(value) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def read_rows(name: str, raw: bytes) -> list[dict]:
    check_file(raw)
    suffix = Path(name).suffix.lower()
    try:
        if suffix == ".csv":
            text = raw.decode("utf-8-sig")
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(io.StringIO(text, newline=""), dialect)
            rows = list(_bounded(reader))
        elif suffix == ".xlsx":
            check_zip(raw)
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            try:
                rows = list(_bounded(book.active.iter_rows(values_only=True)))
            finally:
                book.close()
        else:
            raise AppError("FILE-TYPE", "Import a CSV or XLSX file.")
        if len(rows) < 2:
            raise AppError("FILE-ROWS", "Include a header and at least one data row.")
        keys = [header(k) for k in rows[0]]
        if any(not key for key in keys) or len(keys) != len(set(keys)):
            raise AppError("FILE-HEADERS", "Use unique, nonempty column names.")
        return [
            {
                key: str(value).strip() if value is not None else ""
                for key, value in zip(keys, row, strict=False)
                if key
            }
            for row in rows[1:]
            if any(value is not None and str(value).strip() for value in row)
        ]
    except AppError:
        raise
    except Exception as exc:
        raise AppError(
            "FILE-READ", "Could not read this spreadsheet. Try a UTF-8 CSV export."
        ) from exc


def _bounded(rows):
    for index, row in enumerate(rows):
        if index > MAX_ROWS or len(row) > 150:
            raise AppError("FILE-SIZE", "Imports support up to 10,000 rows and 150 columns.")
        if any(len(str(value or "")) > 60000 for value in row):
            raise AppError("FILE-SIZE", "A cell is too large.")
        yield row


def safe_cell(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, str):
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
        if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r")):
            return "'" + value
    return value


BASE_COLUMNS = [
    'rank', 'fit_score', 'deal_score', 'confidence', 'eligibility', 'next_action', 'title', 'mode',
    'address', 'unit', 'location', 'property_type', 'listing_url', 'inquiry_url', 'broker_url',
    'price_amount', 'currency', 'price_basis', 'area_sqft', 'bedrooms', 'bathrooms', 'lease_type',
    'estimated_monthly_total', 'known_monthly_cost', 'effective_monthly', 'cash_needed', 'known_cash_needed',
    'cost_unknowns', 'cost_assumptions', 'monthly_breakdown', 'cash_breakdown', 'costs',
    'comparable_count', 'comparable_median', 'asking_discount_percent', 'comparison_unit', 'comparison_label',
    'comparables', 'verified_comparable_urls', 'posted_at', 'available_from', 'valid_through',
    'listing_status', 'listing_source', 'source_provider', 'last_checked_at', 'source_check',
    'description', 'features', 'score_breakdown', 'blockers', 'questions', 'due_diligence',
    'research_summary', 'signals', 'outreach_angle', 'risks', 'contact_name', 'contact_role',
    'contact_source', 'email', 'email_status', 'email_source', 'email_checked_at', 'phone',
    'verification_source', 'verification_at', 'research_status', 'status', 'ready_to_tour',
    'tour_blockers', 'ready_to_contact', 'export_blockers', 'review_notes', 'workflow_status',
    'activity_at', 'next_follow_up', 'prior_activity',
]
EXPORT_COLUMNS = BASE_COLUMNS + [
    f'email_{step}_{field}' for step in (1, 2, 3)
    for field in ('subject', 'body', 'delay_days', 'source_urls', 'user_claims')
] + ['client_name', 'reply_email', 'signature', 'source_urls', 'evidence', 'created_at', 'updated_at', 'property_id', 'search_id']


def export_rows(search: dict, opportunities: list[dict]) -> list[dict]:
    result = []
    signature = client_signature(search['client'])
    for opportunity in opportunities:
        row = {key: opportunity.get(key, '') for key in BASE_COLUMNS}
        for key in ('blockers', 'questions', 'tour_blockers', 'export_blockers'):
            row[key] = '; '.join(opportunity.get(key, []))
        estimate, deal = opportunity.get('estimate', {}), opportunity.get('deal', {})
        for target, source in {'estimated_monthly_total': 'monthly_total', 'known_monthly_cost': 'known_monthly', 'effective_monthly': 'effective_monthly', 'cash_needed': 'cash_needed', 'known_cash_needed': 'known_cash_needed', 'cost_unknowns': 'unknowns', 'cost_assumptions': 'assumptions', 'monthly_breakdown': 'monthly_breakdown', 'cash_breakdown': 'cash_breakdown'}.items():
            row[target] = estimate.get(source)
        for target, source in {'comparable_count': 'comparable_count', 'comparable_median': 'median_unit_price', 'asking_discount_percent': 'discount_percent', 'comparison_unit': 'unit', 'comparison_label': 'label'}.items():
            row[target] = deal.get(source)
        for draft in opportunity.get('drafts', []):
            for field in ('subject', 'body', 'delay_days', 'source_urls', 'user_claims'):
                value = draft.get(field, '')
                if field == 'body' and signature:
                    value += '\n\n' + signature
                row[f'email_{draft["step"]}_{field}'] = value
        row.update(client_name=search['client']['name'], reply_email=search['client']['email'], signature=signature)
        for key in ('source_urls', 'evidence', 'created_at', 'updated_at', 'search_id'):
            row[key] = opportunity.get(key, '')
        row['property_id'] = opportunity['id']
        result.append(row)
    return result


def spreadsheet(search: dict, opportunities: list[dict], format: str) -> bytes:
    rows = export_rows(search, opportunities)
    if format == 'csv':
        buffer = io.StringIO(newline='')
        writer = csv.DictWriter(buffer, fieldnames=EXPORT_COLUMNS)
        writer.writeheader()
        writer.writerows({key: safe_cell(row.get(key, '')) for key in EXPORT_COLUMNS} for row in rows)
        return buffer.getvalue().encode('utf-8-sig')
    if format != 'xlsx':
        raise AppError('EXPORT-FORMAT', 'Choose CSV or XLSX.')
    book = Workbook()
    sheet = book.active
    sheet.title = 'Properties'
    sheet.append(EXPORT_COLUMNS)
    for row in rows:
        sheet.append([safe_cell(row.get(key, '')) for key in EXPORT_COLUMNS])
    sheet.freeze_panes = 'H2'
    sources = book.create_sheet('Sources')
    sources.append(['Property ID', 'Property', 'Field', 'Source URL', 'Excerpt', 'Checked at'])
    comps = book.create_sheet('Comparables')
    comps.append(['Property ID', 'Comparable', 'URL', 'Amount', 'Currency', 'Basis', 'Area sqft', 'Asking or closed', 'Observed at', 'Page matched', 'Rationale'])
    for opportunity in opportunities:
        for evidence in opportunity.get('evidence', []):
            sources.append([safe_cell(x) for x in (opportunity['id'], opportunity['title'], evidence['field'], evidence['url'], evidence['quote'], evidence['checked_at'])])
        for comp in opportunity.get('comparables', []):
            comps.append([safe_cell(x) for x in (opportunity['id'], comp['title'], comp['url'], comp['price_amount'], comp['currency'], comp['price_basis'], comp['area_sqft'], comp['basis'], comp['observed_at'], comp['url'] in opportunity.get('verified_comparable_urls', []), comp['rationale'])])
    for tab in (sheet, sources, comps):
        tab.freeze_panes = tab.freeze_panes or 'A2'
        tab.auto_filter.ref = tab.dimensions
        for cell in tab[1]:
            cell.fill = PatternFill('solid', fgColor='1746C7')
            cell.font = Font(color='FFFFFF', bold=True)
            tab.column_dimensions[cell.column_letter].width = 25
        for row in tab.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
    notes = book.create_sheet('Read me')
    for row in [
        ['Search', search['name']],
        ['Data', 'FICTIONAL DEMO — DO NOT CONTACT' if search['demo'] else 'Property research and inquiry drafts. No messages, applications or offers have been sent.'],
        ['Fit', 'Weighted evidence-backed suitability. Unknowns earn no fit points. It is not loan approval, legal eligibility or a purchase recommendation.'],
        ['Confidence', 'Evidence coverage, source matching and cost completeness; not a probability of accuracy.'],
        ['Deal', 'Asking-price comparison requires three eligible recently checked comps. It is not an appraisal, sold-comps valuation, ROI forecast or guarantee of a bargain.'],
        ['Costs', 'Blank total means missing costs. Known costs are a partial subtotal, not an all-in quote. Explicit zero means the user or source stated zero/included.'],
        ['Buying', 'Fixed-rate amortization scenario. Rate, closing, reserves and other assumptions are user inputs, not current loan quotes. Commercial balloon/interest-only structures are not modeled.'],
        ['Leasing', 'Gross base rent and concession-adjusted effective rent are separate. Cash budget includes first-month costs and refundable deposit; payment dates and refund terms need confirmation.'],
        ['Comparisons', 'No automatic currency conversion. Comparable eligibility uses matching mode, currency, named location, property type and size; commercial leases also require matching lease structure.'],
        ['Outreach', 'Email bodies include your signature. Outreach export needs current source checks and your approval, and contains one row per mailbox.'],
        ['Verification', 'Published means the address appeared on a checked source. MX-found confirms domain routing only, not mailbox delivery or identity.'],
        ['Follow-up', 'Days are suggestions from first contact. Stop after a reply, rejection or bounce. Tracking is manual or imported.'],
        ['Privacy', 'This file contains supplied contact details and selected brief facts. Review before sharing.'],
    ]:
        notes.append([safe_cell(x) for x in row])
    notes.column_dimensions['A'].width = 20
    notes.column_dimensions['B'].width = 110
    for row in notes:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
